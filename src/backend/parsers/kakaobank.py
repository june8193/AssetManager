"""카카오뱅크 암호화 엑셀 명세서 파서 모듈."""

import io
import re
from typing import Any, Dict, List, Optional
import msoffcrypto
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from src.backend.parsers.exceptions import (
    ExpenseDecryptionError,
    ExpenseParserError,
    InvalidPasswordError,
)


def parse_kakaobank_excel(
    content: bytes | io.BytesIO, password: str
) -> Dict[str, Any]:
    """카카오뱅크 암호화 엑셀 파일을 복호화하고 표준 거래 내역 형태로 파싱합니다.

    Args:
        content: 엑셀 파일의 원본 바이너리 데이터 또는 BytesIO 스트림.
        password: 엑셀 복호화 비밀번호 (생년월일 6자리 등).

    Returns:
        표준 지출 스키마 딕셔너리:
            - year_month (str): 기준 년월 (예: '2026-08')
            - owner (Optional[str]): 소유주/예금주 성명
            - institution (str): 금융기관명 ('카카오뱅크')
            - account_identifier (Optional[str]): 계좌번호 식별값
            - transactions (List[dict]): 파싱된 거래 내역 목록

    Raises:
        InvalidPasswordError: 비밀번호가 올바르지 않은 경우.
        ExpenseParserError: 파일 손상 또는 엑셀 구조 파싱 실패 시.
    """
    if isinstance(content, bytes):
        input_stream = io.BytesIO(content)
    else:
        input_stream = content
        input_stream.seek(0)

    decrypted_stream = io.BytesIO()

    try:
        office_file = msoffcrypto.OfficeFile(input_stream)
        office_file.load_key(password=password)
        office_file.decrypt(decrypted_stream)
    except msoffcrypto.exceptions.InvalidKeyError as exc:
        raise InvalidPasswordError("카카오뱅크 엑셀 비밀번호가 일치하지 않습니다.") from exc
    except Exception as exc:
        # 암호화되지 않은 일반 엑셀 파일일 수도 있으므로 확인 시도
        input_stream.seek(0)
        try:
            decrypted_stream = io.BytesIO(input_stream.read())
        except Exception:
            raise ExpenseParserError(f"카카오뱅크 파일 복호화 실패: {exc}") from exc

    decrypted_stream.seek(0)
    try:
        workbook = openpyxl.load_workbook(decrypted_stream, data_only=True)
    except Exception as exc:
        raise ExpenseParserError(f"엑셀 워크북을 읽을 수 없습니다: {exc}") from exc

    sheet: Optional[Worksheet] = None
    for name in workbook.sheetnames:
        if "카카오뱅크" in name or "거래내역" in name:
            sheet = workbook[name]
            break
    if sheet is None:
        sheet = workbook.active

    if sheet is None:
        raise ExpenseParserError("카카오뱅크 시트를 찾을 수 없습니다.")

    owner: Optional[str] = None
    account_identifier: Optional[str] = None
    year_month: Optional[str] = None

    header_row_idx: Optional[int] = None
    col_map: Dict[str, int] = {}

    # 상단 메타데이터 및 헤더 행 탐색
    for r in range(1, min(sheet.max_row + 1, 30)):
        for c in range(1, min(sheet.max_column + 1, 15)):
            val = sheet.cell(row=r, column=c).value
            if val is None:
                continue
            str_val = str(val).strip()

            if str_val == "성명":
                owner_cell = sheet.cell(row=r, column=c + 1).value
                if owner_cell:
                    owner = str(owner_cell).strip()
            elif str_val == "계좌번호":
                acc_cell = sheet.cell(row=r, column=c + 1).value
                if acc_cell:
                    account_identifier = str(acc_cell).strip()
            elif str_val == "조회기간":
                period_cell = sheet.cell(row=r, column=c + 1).value
                if period_cell:
                    m = re.search(r"(\d{4})[./-](\d{2})", str(period_cell))
                    if m:
                        year_month = f"{m.group(1)}-{m.group(2)}"
            elif str_val == "거래일시":
                header_row_idx = r
                break

        if header_row_idx is not None:
            # 헤더 컬럼 인덱스 매핑 구성
            for col_idx in range(1, min(sheet.max_column + 1, 15)):
                h_val = sheet.cell(row=header_row_idx, column=col_idx).value
                if h_val:
                    col_map[str(h_val).strip()] = col_idx
            break

    if header_row_idx is None:
        raise ExpenseParserError("카카오뱅크 거래내역 테이블 헤더를 찾을 수 없습니다.")

    date_col = col_map.get("거래일시")
    type_col = col_map.get("구분")
    amount_col = col_map.get("거래금액")
    tx_type_col = col_map.get("거래구분")
    merchant_col = col_map.get("내용")
    memo_col = col_map.get("메모")

    if not date_col or not amount_col:
        raise ExpenseParserError("필수 컬럼(거래일시, 거래금액)을 찾을 수 없습니다.")

    transactions: List[Dict[str, Any]] = []

    for r in range(header_row_idx + 1, sheet.max_row + 1):
        raw_date = sheet.cell(row=r, column=date_col).value
        if not raw_date:
            continue

        raw_date_str = str(raw_date).strip()
        # '2026.08.01 10:28:41' -> '2026-08-01 10:28:41'
        norm_date = re.sub(r"^(\d{4})\.(\d{2})\.(\d{2})", r"\1-\2-\3", raw_date_str)

        # 개별 거래의 year_month 산출
        tx_ym = norm_date[:7]
        if not year_month:
            year_month = tx_ym

        # 금액 추출 및 정규화
        raw_amount = sheet.cell(row=r, column=amount_col).value
        if raw_amount is None:
            continue

        if isinstance(raw_amount, (int, float)):
            amount_val = abs(float(raw_amount))
        else:
            clean_num = str(raw_amount).replace(",", "").strip()
            amount_val = abs(float(clean_num))

        raw_type = (
            str(sheet.cell(row=r, column=type_col).value).strip() if type_col else ""
        )
        tx_type = (
            str(sheet.cell(row=r, column=tx_type_col).value).strip()
            if tx_type_col
            else ""
        )
        raw_merchant = (
            str(sheet.cell(row=r, column=merchant_col).value).strip()
            if merchant_col
            else ""
        )
        raw_memo = (
            str(sheet.cell(row=r, column=memo_col).value).strip()
            if memo_col and sheet.cell(row=r, column=memo_col).value
            else ""
        )

        # 원본 거래구분 결정: '구분'(출금/입금) 우선, 필요시 거래구분(계좌간자동이체 등)
        original_type = raw_type if raw_type else tx_type
        # 메모가 비어있고 거래구분이 있으면 거래구분을 보존
        final_memo = raw_memo if raw_memo else (tx_type if tx_type != original_type else "")

        transactions.append(
            {
                "transaction_date": norm_date,
                "year_month": tx_ym,
                "merchant": raw_merchant,
                "amount": amount_val,
                "original_type": original_type,
                "memo": final_memo,
            }
        )

    return {
        "year_month": year_month or "",
        "owner": owner,
        "institution": "카카오뱅크",
        "account_identifier": account_identifier,
        "transactions": transactions,
    }
